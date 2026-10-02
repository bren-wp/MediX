#!/usr/bin/env python3
"""Discover the public HALMED human-medicine search and export routes."""

from __future__ import annotations

import io
import json
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import load_workbook
from bs4 import BeautifulSoup

BASE_URL = "https://www.halmed.hr/Lijekovi/Baza-lijekova/"


def clean(value: str | None) -> str:
    return " ".join((value or "").split())


def serialize_form(form, page_url: str) -> dict:
    fields = []
    for element in form.find_all(["input", "select", "button", "textarea"]):
        fields.append(
            {
                "tag": element.name,
                "type": element.get("type"),
                "name": element.get("name"),
                "value": element.get("value"),
                "id": element.get("id"),
                "text": clean(element.get_text(" ", strip=True)),
            }
        )

    return {
        "method": (form.get("method") or "get").lower(),
        "action": urljoin(page_url, form.get("action") or page_url),
        "id": form.get("id"),
        "class": form.get("class"),
        "fields": fields,
    }


def export_links(soup: BeautifulSoup, page_url: str) -> list[dict]:
    links = []
    for link in soup.find_all("a", href=True):
        href = urljoin(page_url, link.get("href"))
        text = clean(link.get_text(" ", strip=True))
        lower = (href + " " + text).lower()
        if any(token in lower for token in (".xls", ".xlsx", "excel", "export", "preuz")):
            links.append({"text": text, "href": href})
    return links


def medicine_links(soup: BeautifulSoup, page_url: str) -> list[dict]:
    seen: set[str] = set()
    links: list[dict] = []
    base_path = urlparse(BASE_URL).path.rstrip("/") + "/"
    for link in soup.find_all("a", href=True):
        href = urljoin(page_url, link.get("href"))
        path = urlparse(href).path
        if not path.startswith(base_path) or path.rstrip("/") == base_path.rstrip("/"):
            continue
        if href in seen:
            continue
        seen.add(href)
        links.append(
            {
                "text": clean(link.get_text(" ", strip=True)),
                "href": href,
            }
        )
    return links


def main() -> int:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "MediX-catalog-sync/0.3.1 "
                "(public HALMED human-medicine catalog synchronization; "
                "https://github.com/bren-wp/MediX)"
            )
        }
    )

    response = session.get(BASE_URL, timeout=60)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")
    forms = [serialize_form(form, response.url) for form in soup.find_all("form")]

    search_form = soup.find("form", id="pretrazi_bazu")
    search_report: dict = {}
    if search_form is not None:
        action = urljoin(
            response.url,
            (search_form.get("action") or BASE_URL).split("#", 1)[0],
        )
        # Query one bounded medicine group so discovery cannot be blocked by
        # rendering the entire human-medicine registry in one HTML response.
        # 40 = OTC group in the current public HALMED search form.
        payload = {
            "trazi_baza": "OK",
            "skupine_lijekova": "40",
        }
        result = session.post(
            action,
            data=payload,
            timeout=120,
            allow_redirects=True,
        )
        result.raise_for_status()
        result_soup = BeautifulSoup(result.text, "html.parser")
        meds = medicine_links(result_soup, result.url)

        exports = export_links(result_soup, result.url)
        workbook_report: dict = {}
        if exports:
            xlsx = session.get(exports[0]["href"], timeout=120)
            xlsx.raise_for_status()
            workbook = load_workbook(
                io.BytesIO(xlsx.content),
                read_only=True,
                data_only=True,
            )
            sheet = workbook[workbook.sheetnames[0]]
            sample_rows = []
            for row in sheet.iter_rows(min_row=1, max_row=8, values_only=True):
                sample_rows.append(
                    [None if value is None else str(value) for value in row]
                )
            workbook_report = {
                "content_length": len(xlsx.content),
                "sheet_names": workbook.sheetnames,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "sample_rows": sample_rows,
            }

        search_report = {
            "status": result.status_code,
            "final_url": result.url,
            "content_length": len(result.content),
            "export_links": exports,
            "medicine_link_count": len(meds),
            "sample_medicine_links": meds[:15],
            "contains_results_anchor": bool(
                result_soup.find(id="rezultati")
            ),
            "workbook": workbook_report,
        }

    print(
        json.dumps(
            {
                "landing": {
                    "status": response.status_code,
                    "final_url": response.url,
                    "forms": forms,
                    "export_links": export_links(soup, response.url),
                },
                "sample_group_search": search_report,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
