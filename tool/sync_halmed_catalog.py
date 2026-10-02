#!/usr/bin/env python3
"""Synchronize the public HALMED human-medicine registry for MediX.

HALMED is the medicine-identity source. HZZO and price catalogs may enrich
HALMED records, but may never create medicine identities on their own.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

BASE_URL = "https://www.halmed.hr/Lijekovi/Baza-lijekova/"
SEARCH_FORM_ID = "pretrazi_bazu"
MIN_EXPECTED_RECORDS = 4000

ATC_CATEGORIES = {
    "A": "Probavni sustav i metabolizam",
    "B": "Krv i krvotvorni organi",
    "C": "Srce i krvožilni sustav",
    "D": "Dermatološki lijekovi",
    "G": "Mokraćni i spolni sustav",
    "H": "Hormonski lijekovi",
    "J": "Antiinfektivni lijekovi",
    "L": "Antineoplastici i imunomodulatori",
    "M": "Mišićno-koštani sustav",
    "N": "Živčani sustav",
    "P": "Antiparazitici",
    "R": "Dišni sustav",
    "S": "Osjetila",
    "V": "Razni pripravci (ATK V)",
}

STRENGTH_RE = re.compile(
    r"(\d+(?:[.,]\d+)?\s*(?:mg|g|µg|mcg|mikrograma|ml|mmol|IU|i\.j\.|%|"
    r"jedinica)(?:\s*/\s*\d+(?:[.,]\d+)?\s*(?:ml|g))?)",
    re.IGNORECASE,
)
ATC_TOKEN_RE = re.compile(
    r"^(?:[A-Z]|[A-Z]\d{2}|[A-Z]\d{2}[A-Z]|"
    r"[A-Z]\d{2}[A-Z]{2}|[A-Z]\d{2}[A-Z]{2}\d{2})$"
)
LEGAL_ENTITY_RE = re.compile(
    r"(?:^|\s)(?:d\.?\s*o\.?\s*o\.?|d\.?\s*d\.?|j\.?\s*d\.?\s*o\.?\s*o\.?|"
    r"obrt|ustanova|limited|ltd\.?|gmbh|s\.?a\.?|b\.?v\.?)(?:\s|$)",
    re.IGNORECASE,
)
TECHNICAL_NAMES = {
    "naziv",
    "naziv lijeka",
    "proizvođač",
    "nositelj odobrenja",
    "nije navedeno",
    "nepoznato",
}

REQUIRED_HEADERS = {
    "Naziv",
    "Broj odobrenja",
    "Djelatna tvar",
    "Farmaceutski oblik",
    "Nositelj odobrenja",
    "Način izdavanja",
    "ATK",
}


def clean(value: object | None) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def normalized(value: object | None) -> str:
    return clean(value).casefold()


def stable_id(approval: str, name: str) -> str:
    basis = f"{approval}|{name}".encode("utf-8")
    return "halmed-" + hashlib.sha1(basis).hexdigest()[:24]


def normalize_atc(value: object | None) -> str:
    text = clean(value).upper()
    if text in {"", "-"}:
        return ""
    return "; ".join(
        part.strip()
        for part in re.split(r"[;,]", text)
        if part.strip()
    )


def is_valid_atc(value: object | None) -> bool:
    text = normalize_atc(value)
    if not text:
        return True
    return all(
        ATC_TOKEN_RE.fullmatch(part.strip()) is not None
        for part in text.split(";")
        if part.strip()
    )


def category_for(atc: str) -> str:
    first = normalize_atc(atc)[:1]
    return ATC_CATEGORIES.get(first, "Neklasificirano")


def derive_strength(name: str, composition: str) -> str:
    for value in (name, composition):
        match = STRENGTH_RE.search(value or "")
        if match:
            return clean(match.group(1))
    return ""


def find_export_url(html: str, page_url: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for link in soup.find_all("a", href=True):
        href = urljoin(page_url, link.get("href"))
        text = clean(link.get_text(" ", strip=True)).casefold()
        lower_href = href.casefold()
        if lower_href.endswith((".xls", ".xlsx")):
            return href
        if "spremi rezultate" in text and "xls" in lower_href:
            return href
    raise RuntimeError("HALMED Excel export link was not discovered.")


def medicine_links(html: str, page_url: str) -> dict[str, str]:
    soup = BeautifulSoup(html, "html.parser")
    result: dict[str, str] = {}
    prefix = "/Lijekovi/Baza-lijekova/"
    for link in soup.find_all("a", href=True):
        href = urljoin(page_url, link.get("href"))
        if prefix not in href or href.rstrip("/") == BASE_URL.rstrip("/"):
            continue
        title = clean(link.get_text(" ", strip=True))
        if title:
            result.setdefault(title.casefold(), href)
    return result


def find_header_row(sheet) -> tuple[int, list[str]]:
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=40, values_only=True),
        start=1,
    ):
        headers = [clean(value) for value in row]
        if REQUIRED_HEADERS.issubset(set(headers)):
            return row_number, headers
    raise RuntimeError("HALMED XLSX header row was not recognized.")


def identity_issue(record: dict) -> str | None:
    name = clean(record.get("name"))
    approval = clean(record.get("authorization_number"))
    if not name:
        return "missing_name"
    if not approval:
        return "missing_authorization_number"

    normalized_name = name.casefold()
    if normalized_name in TECHNICAL_NAMES:
        return "technical_placeholder_name"

    holder = normalized(record.get("holder"))
    manufacturer = normalized(record.get("manufacturer"))
    if holder and normalized_name == holder:
        return "name_equals_holder"
    if manufacturer and normalized_name == manufacturer:
        return "name_equals_manufacturer"
    if LEGAL_ENTITY_RE.search(name):
        return "legal_entity_as_name"

    evidence = (
        clean(record.get("active_ingredient")),
        clean(record.get("form")),
        clean(record.get("atc_code")),
        clean(record.get("package")),
        clean(record.get("rx_status")),
    )
    if not any(evidence):
        return "missing_regulatory_metadata"
    return None


def parse_workbook(
    content: bytes,
    detail_links: dict[str, str],
) -> tuple[list[dict], Counter[str]]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
    )
    sheet = workbook[workbook.sheetnames[0]]
    header_row, headers = find_header_row(sheet)
    header_index = {
        header: index
        for index, header in enumerate(headers)
        if header
    }

    def value(row: tuple, header: str) -> str:
        index = header_index.get(header)
        if index is None or index >= len(row):
            return ""
        return clean(row[index])

    records: list[dict] = []
    seen_ids: set[str] = set()
    rejected: Counter[str] = Counter()

    for row in sheet.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    ):
        name = value(row, "Naziv")
        approval = value(row, "Broj odobrenja")
        revoked = value(row, "Datum ukidanja rješenja")

        if not name:
            rejected["missing_name"] += 1
            continue
        if not approval:
            rejected["missing_authorization_number"] += 1
            continue
        if revoked:
            rejected["revoked_authorization"] += 1
            continue

        atc = normalize_atc(value(row, "ATK"))
        composition = value(row, "Sastav")
        record = {
            "id": stable_id(approval, name),
            "name": name,
            "previous_name": value(row, "Raniji naziv"),
            "authorization_number": approval,
            "active_ingredient": value(row, "Djelatna tvar"),
            "composition": composition,
            "form": value(row, "Farmaceutski oblik"),
            "strength": derive_strength(name, composition),
            "package": value(row, "Pakiranje"),
            "manufacturer": value(row, "Proizvođač"),
            "holder": value(row, "Nositelj odobrenja"),
            "decision_date": value(row, "Datum rješenja"),
            "authorization_expiry": value(row, "Rok rješenja"),
            "class": value(row, "Klasa"),
            "urbroj": value(row, "Urbroj"),
            "rx_status": value(row, "Način izdavanja"),
            "prescribing_mode": value(row, "Način propisivanja"),
            "dispensing_place": value(row, "Mjesto izdavanja"),
            "advertising": value(
                row,
                "Način oglašavanja prema stanovništvu",
            ),
            "atc_code": atc,
            "category": category_for(atc),
            "market_status": value(row, "Status lijeka na tržištu"),
            "shortage_status": value(row, "Status nestašice"),
            "official_record_url": detail_links.get(name.casefold()),
        }

        issue = identity_issue(record)
        if issue is not None:
            rejected[issue] += 1
            continue

        record_id = record["id"]
        if record_id in seen_ids:
            rejected["duplicate_id"] += 1
            continue
        seen_ids.add(record_id)
        records.append(record)

    records.sort(
        key=lambda item: (
            item["name"].casefold(),
            item["authorization_number"],
        )
    )
    return records, rejected


def validate(records: list[dict], minimum: int) -> None:
    if len(records) < minimum:
        raise RuntimeError(
            f"Only {len(records)} current HALMED medicines parsed; "
            f"minimum is {minimum}."
        )

    ids = [record["id"] for record in records]
    if len(ids) != len(set(ids)):
        raise RuntimeError("Duplicate HALMED IDs detected.")

    identities = [
        (
            normalized(record.get("authorization_number")),
            normalized(record.get("name")),
        )
        for record in records
    ]
    if len(identities) != len(set(identities)):
        raise RuntimeError("Duplicate HALMED medicine identities detected.")

    invalid_atc = [
        clean(record.get("atc_code"))
        for record in records
        if not is_valid_atc(record.get("atc_code"))
    ]
    if invalid_atc:
        raise RuntimeError(
            "Invalid HALMED ATC values detected: "
            + ", ".join(sorted(set(invalid_atc))[:10])
        )

    invalid_identity = [
        (record.get("name"), identity_issue(record))
        for record in records
        if identity_issue(record) is not None
    ]
    if invalid_identity:
        raise RuntimeError(
            "Invalid medicine identities remained after filtering: "
            + ", ".join(
                f"{name!r} ({reason})"
                for name, reason in invalid_identity[:10]
            )
        )

    names = {normalized(record.get("name")) for record in records}
    if "a1 d.o.o." in names:
        raise RuntimeError("Regression: A1 d.o.o. leaked into medicine names.")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("assets/data/halmed_catalog.json"),
    )
    parser.add_argument(
        "--minimum-records",
        type=int,
        default=MIN_EXPECTED_RECORDS,
    )
    args = parser.parse_args()

    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "MediX-catalog-sync/0.3.1 "
                "(public HALMED human-medicine catalog synchronization; "
                "https://github.com/bren-wp/MediX)"
            )
        }
    )

    landing = session.get(BASE_URL, timeout=60)
    landing.raise_for_status()
    landing_soup = BeautifulSoup(landing.text, "html.parser")
    search_form = landing_soup.find("form", id=SEARCH_FORM_ID)
    if search_form is None:
        raise RuntimeError("HALMED medicine search form was not found.")

    action = urljoin(
        landing.url,
        (search_form.get("action") or BASE_URL).split("#", 1)[0],
    )
    search = session.post(
        action,
        data={"trazi_baza": "OK"},
        timeout=180,
        allow_redirects=True,
    )
    search.raise_for_status()

    export_url = find_export_url(search.text, search.url)
    links = medicine_links(search.text, search.url)

    export = session.get(export_url, timeout=180)
    export.raise_for_status()

    records, rejected = parse_workbook(export.content, links)
    validate(records, args.minimum_records)

    payload = {
        "schema_version": 3,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "name": "HALMED Baza lijekova",
            "url": BASE_URL,
            "search_action": action,
            "export_url": export_url,
            "scope": "public_human_medicine_search_results",
        },
        "quality": {
            "accepted_records": len(records),
            "rejected_records": sum(rejected.values()),
            "rejected_reasons": dict(sorted(rejected.items())),
        },
        "record_count": len(records),
        "records": records,
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(
        f"Wrote {len(records)} current HALMED human medicines "
        f"to {args.output}; rejected={sum(rejected.values())}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
